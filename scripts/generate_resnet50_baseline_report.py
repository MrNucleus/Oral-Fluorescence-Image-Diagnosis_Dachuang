"""Generate the ResNet-50 baseline preliminary report in DOCX and Markdown."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RUNS = (
    PROJECT_ROOT
    / "experiments"
    / "oral_v7_resnet18_baseline_20261026"
    / "code"
    / "outputs"
    / "image_classifier_runs"
)
RESNET50_RUN = RUNS / "resnet50_whole_image_geometry_20261027_pycharm"
RESNET18_RUN = RUNS / "v7_resnet18_original_20261026"
FIGURES = PROJECT_ROOT / "reports" / "figures" / "resnet50_baseline_20261027"
DOCX_OUTPUT = (
    PROJECT_ROOT
    / "docs"
    / "口腔荧光图像AI辅助诊断_ResNet50基线模型初步报告_2026-10-27.docx"
)
MD_OUTPUT = (
    PROJECT_ROOT
    / "docs"
    / "口腔荧光图像AI辅助诊断_ResNet50基线模型初步报告_2026-10-27.md"
)

BLACK = RGBColor(0, 0, 0)
INK = RGBColor(37, 52, 59)
GRAY = RGBColor(89, 104, 112)
LIGHT_GRAY = RGBColor(125, 139, 146)
WHITE = RGBColor(255, 255, 255)
HEADER_FILL = "1F4E79"
ROW_FILL = "EAF2F8"
BORDER = "D9D9D9"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BORDER, size=4):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)
    cannot_split = OxmlElement("w:cantSplit")
    tr_pr.append(cannot_split)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    tr_pr.append(cannot_split)


def add_page_field(paragraph):
    paragraph.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    paragraph.add_run(" 页")


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    for style_name, size, before, after in (
        ("Heading 1", 16, 15, 8),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("口腔荧光图像AI辅助诊断｜ResNet-50基线初步报告")
    set_run_font(run, size=8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)
    for run in footer.runs:
        set_run_font(run, size=8.5, color=GRAY)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        first, rest = text.split(bold_lead, 1)
        run = paragraph.add_run(first + bold_lead)
        set_run_font(run, bold=True)
        if rest:
            run = paragraph.add_run(rest)
            set_run_font(run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.2
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead)
        paragraph.runs[-1].bold = True
        set_run_font(paragraph.runs[-1], bold=True)
        rest = text[len(bold_lead):]
        if rest:
            run = paragraph.add_run(rest)
            set_run_font(run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_caption(doc, text, before_table=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4 if not before_table else 2)
    paragraph.paragraph_format.space_after = Pt(4 if before_table else 7)
    paragraph.paragraph_format.keep_with_next = before_table
    run = paragraph.add_run(text)
    set_run_font(run, size=8.8, color=GRAY)
    return paragraph


def add_table(doc, headers, rows, widths, caption):
    add_caption(doc, caption, before_table=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    repeat_table_header(header_row)
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = header_row.cells[index]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, HEADER_FILL)
        set_cell_margins(cell)
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(header))
        set_run_font(run, size=8.8, bold=True, color=WHITE)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for column_index, (value, width) in enumerate(zip(values, widths)):
            cell = row.cells[column_index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else ROW_FILL)
            set_cell_margins(cell)
            set_cell_border(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(str(value))
            set_run_font(run, size=8.8, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, image_path: Path, caption: str, alt_text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(6.35))
    shape._inline.docPr.set("descr", alt_text)
    add_caption(doc, caption)


def add_cover(doc, training, test):
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)
    run = title.add_run("口腔荧光图像AI辅助诊断\nResNet-50基线模型初步报告")
    set_run_font(run, size=26, bold=True, color=BLACK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("整图二分类训练、冻结测试与基线比较")
    set_run_font(run, size=13.5, color=GRAY)

    meta_lines = [
        "报告性质：初步技术报告",
        "数据版本：oral_binary_v5_yolo_2026-10-18_v1",
        "模型：ImageNet预训练 ResNet-50",
        f"完成轮次：{training['completed_epochs']} / {training['maximum_epochs']}（早停）",
        "报告日期：2026年10月27日",
    ]
    for line in meta_lines:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(line)
        set_run_font(run, size=10.5, color=INK)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(34)
    run = note.add_run(
        "本报告用于阶段验收和组内讨论，不构成临床有效性结论。"
    )
    set_run_font(run, size=9.5, color=GRAY)
    doc.add_page_break()


def build_docx(training_record, test_record):
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "口腔荧光图像AI辅助诊断ResNet-50基线模型初步报告"
    doc.core_properties.subject = "整图二分类基线训练、冻结测试与比较"
    doc.core_properties.author = "项目模型训练组"
    doc.core_properties.keywords = "口腔荧光, ResNet-50, 整图分类, 基线模型"

    training = training_record["training"]
    test = training_record["frozen_test"]
    resnet18_test = read_json(RESNET18_RUN / "training_record_20261027.json")["frozen_test"]
    stratified = read_json(
        RESNET50_RUN / "final_test_20261027_1345" / "stratified_metrics.json"
    )
    rows = stratified["rows"]
    site_rows = [row for row in rows if row["dimension"] == "site"]
    modality_rows = [row for row in rows if row["dimension"] == "modality"]

    add_cover(doc, training, test)

    doc.add_heading("一、任务结论", level=1)
    add_body(
        doc,
        "本报告对应模型开发第三阶段第11至12周的基线模型任务：实现ResNet-50基线，"
        "完成可复现训练流程，并获得初步实验结果。当前三项任务均已形成可核验产物，"
        "但该模型仍是阶段基线，不是最终论文模型，也不能作为临床诊断依据。",
    )
    add_table(
        doc,
        ["任务项", "状态", "当前证据"],
        [
            ["实现 ResNet-50 基线", "已完成", "第35轮最佳权重 best.pt"],
            ["完成训练流程代码", "已完成", "PyCharm可运行配置、断点和历史记录"],
            ["获得初步实验结果", "已完成", "验证与冻结测试指标、部位和模态分层"],
            ["最终模型与临床验证", "未完成", "仍需优化、患者关系核验和外部队列"],
        ],
        [2.5, 1.25, 2.9],
        "表1  第11至12周基线任务完成情况",
    )
    add_body(
        doc,
        f"训练在 {training['completed_epochs']} 个 epoch 后早停，最佳点为第 "
        f"{training['best_epoch']} 轮。最佳验证 AUROC 为 {training['best_validation_auroc']:.4f}，"
        f"冻结测试 AUROC 为 {test['auroc']:.4f}，二者存在明显差距。"
        "因此，本报告的核心判断是：训练流程和基线实现已完成，性能尚不足以支持最终模型结论。",
    )

    doc.add_heading("二、数据与评价口径", level=1)
    add_body(
        doc,
        "本轮沿用交接包中的冻结整图数据，不读取XML框、不裁剪ROI，也不使用框级标签。"
        "train、val、test 的来源组划分保持不变；训练和选模仅使用 train/val，test 在权重和阈值冻结后评估一次。",
    )
    add_table(
        doc,
        ["划分", "图片数", "来源组", "正常候选", "病变候选"],
        [
            ["Train", "893", "235", "332", "561"],
            ["Validation", "112", "16", "43", "69"],
            ["Test", "112", "17", "43", "69"],
        ],
        [1.5, 1.1, 1.1, 1.2, 1.2],
        "表2  冻结数据划分与候选标签分布",
    )
    add_body(
        doc,
        "类别映射固定为 0 = normal_candidate、1 = lesion_candidate，正类概率为 "
        "P(lesion_candidate = 1)，判定阈值固定为 0.5。标签仍是候选标签，未经医生逐图确诊。"
        "数据按来源文件夹组隔离，但真实患者关系尚未核验，因此不能把结果表述为患者独立验证。",
    )
    add_body(
        doc,
        "数据身份已锁定：dataset.json 的 SHA-256 为 "
        "5dfb816e0c1f36c418c7e39b2e279c71128f69b2e70eafba192f0279edda4289；"
        "匿名部位表的 SHA-256 为 "
        "5e3e2e67b1df8067bb59ff5f9fd1b30f8045d07ac681606e400e52c466c8af09。",
    )

    doc.add_heading("三、训练流程实现", level=1)
    add_body(
        doc,
        "训练入口是 experiments/oral_v7_resnet18_baseline_20261026/code/scripts/"
        "train_image_classifier.py，使用 PyCharm 运行配置 ResNet50 Baseline (Whole Image)。"
        "模型采用 ImageNet 预训练 ResNet-50，分类头替换为二分类线性层。",
    )
    add_table(
        doc,
        ["项目", "设置"],
        [
            ["输入", "完整RGB图像，640像素，等比缩放并补边"],
            ["训练增强", "geometry v1：水平翻转、轻度旋转、曝光微调"],
            ["优化器", "AdamW，lr=0.001，weight decay=0.0001"],
            ["batch / workers", "8 / 4"],
            ["随机种子", "42"],
            ["选模", "最低整体验证交叉熵"],
            ["早停", "patience=20，最多100 epoch"],
            ["GPU", "RTX 5060 Laptop GPU，CUDA 12.8"],
        ],
        [1.6, 4.7],
        "表3  ResNet-50基线训练配置",
    )
    add_body(
        doc,
        "本机总显存约 7.96 GiB，ResNet-50 在 640 输入下使用 batch 8，避免 OOM；"
        "ResNet18 基线使用 batch 16。该差异会影响训练动态，比较两个模型时需要保留这一限制。"
        "每个 epoch 结束后保存完整 last.pt，最佳验证损失对应的权重另存为 best.pt。",
    )

    add_figure(
        doc,
        FIGURES / "training_curves.png",
        "图1  ResNet-50训练与验证损失、AUROC和平衡准确率。虚线为第35轮最佳点。",
        "ResNet-50 training and validation curves across 55 epochs",
    )

    doc.add_heading("四、训练结果", level=1)
    add_body(
        doc,
        f"训练从 2026-10-27 11:30:52 持续到 13:30:46，共完成 {training['completed_epochs']} 个 epoch。"
        f"第 {training['best_epoch']} 轮验证损失最低；此后连续20轮未改善，训练自动停止。",
    )
    add_table(
        doc,
        ["指标", "第35轮验证结果"],
        [
            ["Validation loss", f"{training['best_validation_loss']:.4f}"],
            ["AUROC", f"{training['best_validation_auroc']:.4f}"],
            ["Accuracy", f"{training['best_validation_accuracy']:.4f}"],
            ["Sensitivity", f"{training['best_validation_sensitivity']:.4f}"],
            ["Specificity", f"{training['best_validation_specificity']:.4f}"],
            ["Balanced accuracy", f"{training['best_validation_balanced_accuracy']:.4f}"],
            ["Confusion matrix", "[[39, 4], [14, 55]]"],
        ],
        [2.2, 3.0],
        "表4  ResNet-50最佳验证结果",
    )
    add_body(
        doc,
        "验证集表现较为均衡：特异度达到0.9070，敏感度为0.7971。训练后期训练损失继续下降，"
        "但验证损失波动明显，说明模型容量和样本量之间的张力已经出现。",
    )

    doc.add_heading("五、冻结测试与分层结果", level=1)
    add_body(
        doc,
        "权重、数据和阈值冻结后，对112张测试图片执行一次评估。测试集来自17个来源组，"
        "正常候选43张，病变候选69张。",
    )
    add_table(
        doc,
        ["测试指标", "结果"],
        [
            ["AUROC", f"{test['auroc']:.4f}"],
            ["Accuracy", f"{test['accuracy']:.4f}"],
            ["Sensitivity", f"{test['sensitivity']:.4f}"],
            ["Specificity", f"{test['specificity']:.4f}"],
            ["Balanced accuracy", f"{test['balanced_accuracy']:.4f}"],
            ["Confusion matrix", "[[25, 18], [14, 55]]"],
        ],
        [2.2, 3.0],
        "表5  ResNet-50冻结测试结果",
    )
    add_body(
        doc,
        "测试 AUROC 为0.6943，明显低于验证集0.9279。验证与测试之间的落差是本轮最重要的风险信号，"
        "可能与样本量、来源组差异、候选标签噪声或模型选择波动有关，不能仅用测试分数解释。",
    )

    site_table_rows = []
    cn_names = {
        "tongue_floor_broad": "舌/口底",
        "gingivopalatal_broad": "牙龈/腭",
        "buccal_labial_vestibular_broad": "颊/唇/前庭",
    }
    for row in sorted(site_rows, key=lambda item: item["images"], reverse=True):
        site_table_rows.append(
            [
                cn_names[row["group"]],
                row["images"],
                row["source_groups"],
                f'{row["auroc"]:.4f}',
                f'{row["sensitivity"]:.4f}',
                f'{row["specificity"]:.4f}',
                f'{row["balanced_accuracy"]:.4f}',
            ]
        )
    add_table(
        doc,
        ["部位组", "n", "来源组", "AUROC", "敏感度", "特异度", "平衡准确率"],
        site_table_rows,
        [1.55, 0.55, 0.7, 0.85, 0.9, 0.9, 1.0],
        "表6  冻结测试的部位分层结果",
    )

    modality_table_rows = []
    for row in modality_rows:
        modality_table_rows.append(
            [
                "荧光" if row["group"] == "fluorescence" else "白光",
                row["images"],
                row["source_groups"],
                f'{row["auroc"]:.4f}',
                f'{row["sensitivity"]:.4f}',
                f'{row["specificity"]:.4f}',
                f'{row["balanced_accuracy"]:.4f}',
            ]
        )
    add_table(
        doc,
        ["模态", "n", "来源组", "AUROC", "敏感度", "特异度", "平衡准确率"],
        modality_table_rows,
        [1.35, 0.55, 0.7, 0.85, 0.9, 0.9, 1.0],
        "表7  冻结测试的模态分层结果",
    )
    add_figure(
        doc,
        FIGURES / "stratified_auroc.png",
        "图2  ResNet-50冻结测试的部位与模态AUROC。",
        "ResNet-50 frozen-test AUROC by anatomical site and image modality",
    )
    add_body(
        doc,
        "白光子组 AUROC 为0.6117，荧光子组为0.8529。该差异只能作为描述性观察："
        "子组来源组数量少，且没有预先规定置信区间，不能据此选择模态或宣称某一模态更可靠。",
    )

    doc.add_heading("六、与ResNet18基线的描述性比较", level=1)
    add_table(
        doc,
        ["模型", "测试AUROC", "敏感度", "特异度", "平衡准确率"],
        [
            [
                "ResNet18",
                f'{resnet18_test["auroc"]:.4f}',
                f'{resnet18_test["sensitivity"]:.4f}',
                f'{resnet18_test["specificity"]:.4f}',
                f'{resnet18_test["balanced_accuracy"]:.4f}',
            ],
            [
                "ResNet50",
                f'{test["auroc"]:.4f}',
                f'{test["sensitivity"]:.4f}',
                f'{test["specificity"]:.4f}',
                f'{test["balanced_accuracy"]:.4f}',
            ],
        ],
        [1.4, 1.15, 1.1, 1.1, 1.35],
        "表8  两个整图基线在同一测试集上的描述性比较",
    )
    add_figure(
        doc,
        FIGURES / "test_comparison.png",
        "图3  ResNet18与ResNet50在冻结测试集上的指标比较。",
        "Frozen test comparison between ResNet18 and ResNet50",
    )
    add_body(
        doc,
        "ResNet50把特异度从0.3488提高到0.5814，但敏感度从0.9855降到0.7971，"
        "AUROC也从0.7988降到0.6943。两个模型在不同指标上各有取舍，当前没有足够证据说明"
        "哪一个应作为最终模型；应继续以验证集为主要决策依据。",
    )

    doc.add_heading("七、问题与局限", level=1)
    for item in [
        "验证与测试差距：ResNet50验证AUROC 0.9279、测试AUROC 0.6943，提示泛化不稳定。",
        "样本量有限：测试集仅112张、17个来源组，子组结果波动较大。",
        "标签边界：候选标签未经医生逐图确诊，不能把模型输出等同于临床诊断。",
        "分组边界：来源文件夹组不等于真实患者，患者独立评价尚未完成。",
        "测试复用：同一测试集已用于ResNet18报告，本轮结果不能视为完全未触碰的外部验证。",
        "训练条件差异：ResNet50受8 GB显存限制使用batch 8，ResNet18使用batch 16。",
        "可解释性缺位：Grad-CAM尚未实现，当前无法判断模型关注区域是否与病灶区域一致。",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    doc.add_heading("八、结论与下一步", level=1)
    add_body(
        doc,
        "ResNet-50基线模型已经完成实现、训练和初步评价，满足第11至12周任务的工程要求，"
        "但性能与稳定性尚不足以进入最终模型阶段。下一步应继续保持测试集封存，"
        "把类别加权、阈值校准、单模态和增强策略作为验证集上的候选改进，不按测试结果逐次挑方案。",
    )
    for item in [
        "完成患者级关系核验，明确当前来源组划分的独立性边界。",
        "在验证集上预注册并比较类别加权、损失函数、阈值校准和单模态方案。",
        "统一比较时固定数据、split、输入尺寸、增强与训练预算，并报告多随机种子波动。",
        "确定候选模型后实现Grad-CAM，检查关注区域与历史标注框的一致性。",
        "在新的独立保留集或外部数据上完成最终验证，再形成论文主结果。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("附录  可复现产物与审计信息", level=1)
    add_table(
        doc,
        ["产物", "路径或标识"],
        [
            ["最佳权重", "resnet50_whole_image_geometry_20261027_pycharm/best.pt"],
            ["训练历史", "history.json（55个epoch）"],
            ["训练总记录", "training_record_20261027.json"],
            ["冻结协议", "final_test_protocol_20261027.json"],
            ["测试指标", "final_test_20261027_1345/test_metrics.json"],
            ["部位/模态分层", "final_test_20261027_1345/stratified_metrics.csv"],
            ["最佳权重SHA-256", "d08425df3b0b2906ecb7174f2aed098b6daa114165e4d7de0589b43eea31b3e4"],
            ["PyCharm配置", ".idea/runConfigurations/ResNet50_Baseline.xml"],
        ],
        [1.65, 4.65],
        "表9  报告依据与可复现产物",
    )
    add_body(
        doc,
        "本报告中的数值来自冻结训练记录和测试输出，不补写未运行的实验。"
        "如果后续更改数据、模型、阈值或评测口径，应新建实验版本并重新记录，"
        "不能覆盖本报告所对应的冻结结果。",
    )
    return doc


def build_markdown(training_record, test_record):
    training = training_record["training"]
    test = training_record["frozen_test"]
    resnet18_test = read_json(RESNET18_RUN / "training_record_20261027.json")["frozen_test"]
    stratified = read_json(
        RESNET50_RUN / "final_test_20261027_1345" / "stratified_metrics.json"
    )
    rows = stratified["rows"]
    site_rows = [row for row in rows if row["dimension"] == "site"]
    modality_rows = [row for row in rows if row["dimension"] == "modality"]
    cn_names = {
        "tongue_floor_broad": "舌/口底",
        "gingivopalatal_broad": "牙龈/腭",
        "buccal_labial_vestibular_broad": "颊/唇/前庭",
    }

    lines = [
        "# 口腔荧光图像AI辅助诊断 ResNet-50基线模型初步报告",
        "",
        "**报告性质：** 初步技术报告，非最终论文或临床验证",
        "",
        "**数据版本：** `oral_binary_v5_yolo_2026-10-18_v1`  ",
        "**模型：** ImageNet预训练 ResNet-50  ",
        f"**完成轮次：** {training['completed_epochs']} / {training['maximum_epochs']}（早停）  ",
        "**报告日期：** 2026年10月27日",
        "",
        "## 一、任务结论",
        "",
        "本报告对应模型开发第三阶段第11至12周任务：实现ResNet-50基线、完成训练流程代码并形成初步结果。三项工程任务均已完成，但模型仍是阶段基线，不能作为临床诊断依据。",
        "",
        "| 任务项 | 状态 | 当前证据 |",
        "|---|---|---|",
        "| 实现ResNet-50基线 | 已完成 | 第35轮最佳权重`best.pt` |",
        "| 完成训练流程代码 | 已完成 | PyCharm运行配置、断点和历史记录 |",
        "| 获得初步实验结果 | 已完成 | 验证、冻结测试及分层指标 |",
        "| 最终模型与临床验证 | 未完成 | 仍需优化、患者关系核验和外部验证 |",
        "",
        f"训练在第{training['best_epoch']}轮取得最佳验证损失{training['best_validation_loss']:.4f}，"
        f"验证AUROC为{training['best_validation_auroc']:.4f}；冻结测试AUROC为{test['auroc']:.4f}。"
        "验证与测试差距明显，因此核心结论是流程和基线已完成，但性能尚不足以支持最终模型结论。",
        "",
        "## 二、数据与评价口径",
        "",
        "本轮沿用冻结整图数据，不读取XML框、不裁剪ROI，也不使用框级标签。train/val/test的来源组划分保持不变；测试集在权重和阈值冻结后只评估一次。",
        "",
        "| 划分 | 图片数 | 来源组 | 正常候选 | 病变候选 |",
        "|---|---:|---:|---:|---:|",
        "| Train | 893 | 235 | 332 | 561 |",
        "| Validation | 112 | 16 | 43 | 69 |",
        "| Test | 112 | 17 | 43 | 69 |",
        "",
        "- `dataset.json` SHA-256：`5dfb816e0c1f36c418c7e39b2e279c71128f69b2e70eafba192f0279edda4289`",
        "- `image_sites.json` SHA-256：`5e3e2e67b1df8067bb59ff5f9fd1b30f8045d07ac681606e400e52c466c8af09`",
        "- 类别：`0 = normal_candidate`，`1 = lesion_candidate`；固定阈值0.5。",
        "- 标签为候选标签，未经医生逐图确诊；来源组不等于已确认患者。",
        "",
        "## 三、训练流程实现",
        "",
        "训练入口为 `experiments/oral_v7_resnet18_baseline_20261026/code/scripts/train_image_classifier.py`，"
        "PyCharm运行配置为 `ResNet50 Baseline (Whole Image)`。",
        "",
        "| 项目 | 设置 |",
        "|---|---|",
        "| 输入 | 完整RGB，640像素，等比缩放并补边 |",
        "| 增强 | geometry v1 |",
        "| 优化器 | AdamW，lr=0.001，weight decay=0.0001 |",
        "| batch / workers | 8 / 4 |",
        "| 随机种子 | 42 |",
        "| 选模 | 最低整体验证交叉熵 |",
        "| 早停 | patience=20，最多100 epoch |",
        "| GPU | RTX 5060 Laptop GPU，CUDA 12.8 |",
        "",
        "本机总显存约7.96 GiB，ResNet-50使用batch 8避免OOM；ResNet18基线使用batch 16。"
        "每个epoch保存完整`last.pt`，最佳验证损失权重另存为`best.pt`。",
        "",
        "![ResNet-50训练曲线](../reports/figures/resnet50_baseline_20261027/training_curves.png)",
        "",
        "## 四、训练结果",
        "",
        f"训练从2026-10-27 11:30:52持续到13:30:46，共完成{training['completed_epochs']}个epoch。"
        f"第{training['best_epoch']}轮验证损失最低，之后连续20轮未改善并早停。",
        "",
        "| 指标 | 第35轮验证结果 |",
        "|---|---:|",
        f"| Validation loss | {training['best_validation_loss']:.4f} |",
        f"| AUROC | {training['best_validation_auroc']:.4f} |",
        f"| Accuracy | {training['best_validation_accuracy']:.4f} |",
        f"| Sensitivity | {training['best_validation_sensitivity']:.4f} |",
        f"| Specificity | {training['best_validation_specificity']:.4f} |",
        f"| Balanced accuracy | {training['best_validation_balanced_accuracy']:.4f} |",
        "| Confusion matrix | `[[39,4],[14,55]]` |",
        "",
        "## 五、冻结测试与分层结果",
        "",
        "| 测试指标 | 结果 |",
        "|---|---:|",
        f"| AUROC | {test['auroc']:.4f} |",
        f"| Accuracy | {test['accuracy']:.4f} |",
        f"| Sensitivity | {test['sensitivity']:.4f} |",
        f"| Specificity | {test['specificity']:.4f} |",
        f"| Balanced accuracy | {test['balanced_accuracy']:.4f} |",
        "| Confusion matrix | `[[25,18],[14,55]]` |",
        "",
        "| 部位组 | n | 来源组 | AUROC | 敏感度 | 特异度 | 平衡准确率 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for row in sorted(site_rows, key=lambda item: item["images"], reverse=True):
        lines.append(
            f'| {cn_names[row["group"]]} | {row["images"]} | {row["source_groups"]} | '
            f'{row["auroc"]:.4f} | {row["sensitivity"]:.4f} | '
            f'{row["specificity"]:.4f} | {row["balanced_accuracy"]:.4f} |'
        )
    lines.extend(
        [
            "",
            "| 模态 | n | 来源组 | AUROC | 敏感度 | 特异度 | 平衡准确率 |",
            "|---|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for row in modality_rows:
        lines.append(
            f'| {"荧光" if row["group"] == "fluorescence" else "白光"} | '
            f'{row["images"]} | {row["source_groups"]} | {row["auroc"]:.4f} | '
            f'{row["sensitivity"]:.4f} | {row["specificity"]:.4f} | '
            f'{row["balanced_accuracy"]:.4f} |'
        )
    lines.extend(
        [
            "",
            "![ResNet-50分层AUROC](../reports/figures/resnet50_baseline_20261027/stratified_auroc.png)",
            "",
            "## 六、与ResNet18基线的描述性比较",
            "",
            "| 模型 | 测试AUROC | 敏感度 | 特异度 | 平衡准确率 |",
            "|---|---:|---:|---:|---:|",
            f'| ResNet18 | {resnet18_test["auroc"]:.4f} | {resnet18_test["sensitivity"]:.4f} | '
            f'{resnet18_test["specificity"]:.4f} | {resnet18_test["balanced_accuracy"]:.4f} |',
            f'| ResNet50 | {test["auroc"]:.4f} | {test["sensitivity"]:.4f} | '
            f'{test["specificity"]:.4f} | {test["balanced_accuracy"]:.4f} |',
            "",
            "![基线测试比较](../reports/figures/resnet50_baseline_20261027/test_comparison.png)",
            "",
            "ResNet50提高了特异度，但敏感度和AUROC低于ResNet18。两个模型在不同指标上各有取舍，"
            "当前没有足够证据说明哪一个应作为最终模型，后续仍应以验证集为主要决策依据。",
            "",
            "## 七、问题与局限",
            "",
            "- 验证AUROC 0.9279与测试AUROC 0.6943差距明显，提示泛化不稳定。",
            "- 测试集仅112张、17个来源组，子组结果波动较大。",
            "- 候选标签未经医生逐图确诊，不能等同于临床诊断。",
            "- 来源文件夹组不等于真实患者，患者独立评价尚未完成。",
            "- 同一测试集已用于ResNet18报告，本轮不能视为完全未触碰的外部验证。",
            "- ResNet50受显存限制使用batch 8，与ResNet18的batch 16不同。",
            "- Grad-CAM尚未实现，当前无法验证模型关注区域。",
            "",
            "## 八、结论与下一步",
            "",
            "ResNet-50基线已完成实现、训练和初步评价，满足第11至12周任务的工程要求，"
            "但性能与稳定性尚不足以进入最终模型阶段。",
            "",
            "- 完成患者级关系核验，明确当前来源组划分的独立性边界。",
            "- 在验证集上预注册并比较类别加权、损失函数、阈值校准和单模态方案。",
            "- 固定数据、split、输入尺寸、增强和训练预算，并报告多随机种子的波动。",
            "- 确定候选模型后实现Grad-CAM并检查关注区域一致性。",
            "- 在新的独立保留集或外部数据上完成最终验证。",
            "",
            "## 附录：报告依据",
            "",
            "- 训练总记录：`training_record_20261027.json`",
            "- 冻结协议：`final_test_protocol_20261027.json`",
            "- 测试指标：`final_test_20261027_1345/test_metrics.json`",
            "- 分层表：`final_test_20261027_1345/stratified_metrics.csv`",
            "- 最佳权重SHA-256：`d08425df3b0b2906ecb7174f2aed098b6daa114165e4d7de0589b43eea31b3e4`",
            "",
        ]
    )
    MD_OUTPUT.write_text("\n".join(lines), encoding="utf-8")


def main():
    training_record = read_json(RESNET50_RUN / "training_record_20261027.json")
    test_record = read_json(
        RESNET50_RUN / "final_test_20261027_1345" / "test_metrics.json"
    )
    for figure in (
        FIGURES / "training_curves.png",
        FIGURES / "test_comparison.png",
        FIGURES / "stratified_auroc.png",
    ):
        if not figure.is_file():
            raise FileNotFoundError(figure)
    build_markdown(training_record, test_record)
    doc = build_docx(training_record, test_record)
    doc.save(DOCX_OUTPUT)
    print(json.dumps({"docx": str(DOCX_OUTPUT), "markdown": str(MD_OUTPUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
