"""Generate the current project progress report as a formatted DOCX."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = (
    PROJECT_ROOT
    / "docs"
    / "口腔荧光图像AI辅助诊断_阶段进展报告_2026-10-19_严子宁_语义确认版.docx"
)
TRAINING_FIGURE = (
    PROJECT_ROOT
    / "experiments"
    / "binary_resnet18_annotated_seed43"
    / "training_curves.png"
)
TEST_FIGURE = (
    PROJECT_ROOT
    / "experiments"
    / "binary_resnet18_annotated_seed43"
    / "test_overview.png"
)

NAVY = RGBColor(31, 78, 121)
DARK = RGBColor(45, 48, 51)
GRAY = RGBColor(95, 99, 104)
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F3F4F6"
PALE_ORANGE = "FCE4D6"


def set_run_font(run, name: str = "宋体", size: float | None = None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.35):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, bottom=90, left=110, right=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C9D6", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, "微软雅黑", 16 if level == 1 else 12.5, True, NAVY if level == 1 else DARK)
    set_paragraph_spacing(paragraph, before=14 if level == 1 else 9, after=6)
    return paragraph


def add_body(doc: Document, text: str, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph)
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, "宋体", 10.5, True, DARK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, "宋体", 10.5, False, DARK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, "宋体", 10.5, False, DARK)
    return paragraph


def add_bullet(doc: Document, text: str, level: int = 0):
    style = "List Bullet 2" if level else "List Bullet"
    paragraph = doc.add_paragraph(style=style)
    set_paragraph_spacing(paragraph, after=3, line=1.25)
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 10.5, False, DARK)
    return paragraph


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    for index, text in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        set_run_font(run, "微软雅黑", 9.5, True, NAVY)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, text in enumerate(row):
            cell = cells[column_index]
            set_cell_margins(cell)
            if row_index % 2 == 1:
                shade_cell(cell, LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(str(text))
            set_run_font(run, "宋体", 9.2, False, DARK)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="AFC4D6", size="8")
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, 150, 150, 180, 180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    set_run_font(run, "微软雅黑", 10.5, True, NAVY)
    paragraph2 = cell.add_paragraph()
    paragraph2.paragraph_format.space_after = Pt(0)
    paragraph2.paragraph_format.line_spacing = 1.3
    run2 = paragraph2.add_run(text)
    set_run_font(run2, "宋体", 10, False, DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, image_path: Path, caption: str, width_cm: float = 15.4):
    if not image_path.exists():
        add_callout(doc, "图片缺失", f"未找到：{image_path}", PALE_ORANGE)
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(caption_paragraph, after=8)
    run = caption_paragraph.add_run(caption)
    set_run_font(run, "宋体", 9, False, GRAY)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, "宋体", 9, False, GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.line_spacing = 1.35

    for level, size in ((1, 16), (2, 12.5), (3, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY if level == 1 else DARK

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("口腔荧光图像AI辅助诊断项目｜阶段进展报告")
    set_run_font(header_run, "微软雅黑", 8.5, False, GRAY)
    add_page_number(section.footer.paragraphs[0])

    props = doc.core_properties
    props.title = "口腔荧光图像AI辅助诊断项目阶段进展报告"
    props.subject = "项目阶段进展总结"
    props.author = "严子宁"
    props.comments = "综合数据处理工作报告与当前模型实验结果整理。"


def build_report():
    doc = Document()
    configure_document(doc)

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("口腔荧光图像AI辅助诊断项目\n阶段进展报告")
    set_run_font(title_run, "微软雅黑", 26, True, NAVY)
    title.paragraph_format.space_after = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("数据处理、二分类数据包与模型探索进展")
    set_run_font(subtitle_run, "微软雅黑", 13.5, False, GRAY)

    for _ in range(5):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run("汇报人：严子宁\n2026年10月19日")
    set_run_font(meta_run, "宋体", 12, False, DARK)
    meta.paragraph_format.line_spacing = 1.8
    doc.add_page_break()

    # Summary
    add_heading(doc, "摘要", 1)
    add_callout(
        doc,
        "当前结论",
        "新版数据已经完成候选整理、二分类框级口径梳理、固定来源组划分和在线增强代码验证；"
        "但新版临床图像尚未完成正式模型训练。当前仓库中的模型结果来自旧版590张数据，"
        "只能作为方法探索，不能与新版1117张候选数据直接比较。",
    )
    add_body(
        doc,
        "本阶段工作可以概括为三条主线。第一，完成旧版与新版原始数据的来源、文件、标注、"
        "方向、隐私遮挡和登记表对应关系核查。第二，将任务从证据不足的五分类尝试收敛为"
        "以检测框为单位的 l/n 二分类准备，并形成可交付的数据包。第三，在旧版数据上建立"
        "二分类训练、阈值选择、患者级评价和过程可视化流程，同时按指导书实现了 ResNet-50 "
        "与 EfficientNet-B3 的三阶段训练框架。",
    )
    add_body(
        doc,
        "当前最需要避免的错误是把“数据整理完成”“代码通过测试”“模型完成训练”和“模型具备"
        "临床价值”混为一谈。本报告分别说明已完成、已验证、待确认和待训练的内容。",
    )

    add_heading(doc, "一、阶段工作总览", 1)
    add_table(
        doc,
        ["工作模块", "当前状态", "结论与边界"],
        [
            ["原始数据与登记信息梳理", "已完成主体工作", "旧版与新版分别统计，不混合比较"],
            ["新版候选数据包", "已形成", "1117张图片、1200个框"],
            ["固定来源组划分", "已完成", "893/112/112，未完成患者级独立性确认"],
            ["YOLO格式转换", "已完成", "图片、XML、TXT逐项对应，含校验记录"],
            ["在线增强代码", "首版完成", "11项合成测试通过，真实训练仍需联调"],
            ["新版临床图像正式训练", "尚未完成", "当前无可报告的新版模型指标"],
            ["旧版590张二分类探索", "已完成多轮", "最佳患者AUROC约0.72，仍不稳定"],
            ["指导书模型参数落地", "框架完成、冒烟通过", "ResNet-50/EfficientNet-B3正式训练待启动"],
        ],
        widths=[4.2, 4.0, 8.2],
    )

    add_heading(doc, "二、数据与标注整理", 1)
    add_heading(doc, "2.1 数据版本必须分开记录", 2)
    add_body(
        doc,
        "旧版数据包含84个目录、590张图片和510份配对XML。新版数据增加到277个目录、"
        "1410张图片和1387份XML，其中1386对能够对应。新版配对率提高不能用于说明旧版"
        "缺标问题已经解决，两者必须作为不同版本分别保留。",
    )
    add_table(
        doc,
        ["阶段", "图片数", "主要统计范围"],
        [
            ["旧版原始资料", "590", "第一轮清点，510对图像与XML"],
            ["新版原始资料", "1410", "1387份XML，1386对可对应"],
            ["严格病例说明筛选", "1017", "按当时四条规则及明确例外保留"],
            ["二分类扩展候选", "1251", "通过结构条件，含1356个框"],
            ["框级筛选后候选", "1117", "至少保留一个可用框，共1200个框"],
        ],
        widths=[4.8, 3.0, 9.0],
    )

    add_heading(doc, "2.2 文件与隐私处理", 2)
    for item in [
        "检查图片、XML、尺寸、框数和边界，识别旧版MPO容器与EXIF方向问题。",
        "工作副本统一解码、去除EXIF、更新XML路径，并保留原文件。",
        "遮挡已知设备信息栏，同时标记遮挡区域与检测框相交的图片。",
        "脱敏交接表不保留完整病例原话和身份字段，私密证据与训练包分开管理。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2.3 五分类与二分类的关系", 2)
    add_body(
        doc,
        "指导书采用 Normal、OLK、OLP、OSCC、Other 五分类，但现有XML中的 l/n "
        "只表示病变区域与正常区域。病例文字中的“癌”、异常增生、细胞学阴性等描述，"
        "不能直接映射为五分类金标准。当前最稳妥的任务是先将每个检测框作为单位，"
        "开展 l/n 二分类准备，五分类等待医生确认的逐图或逐框诊断标签。",
    )

    add_heading(doc, "三、二分类数据包与固定划分", 1)
    add_body(
        doc,
        "当前二分类口径以框为单位：l 为阳性病例，编码1；n 为阴性病例，编码0。"
        "n框本身具有坐标，不等同于整图背景；缺XML或空标注不能直接当作n。"
        "在1251张、1356个框的候选上，排除极小框和低清晰度框后，保留1200个框，"
        "对应1117张图片，其中荧光501张、白光616张。",
    )
    add_table(
        doc,
        ["划分", "图片数", "框数", "来源文件夹组数"],
        [
            ["训练集", "893", "940", "235"],
            ["验证集", "112", "139", "16"],
            ["测试集", "112", "121", "17"],
            ["合计", "1117", "1200", "268"],
        ],
        widths=[4.2, 3.4, 3.4, 5.0],
    )
    add_body(
        doc,
        "这套划分按来源文件夹分组，检查未发现同一来源组或同一图片ID跨集合。"
        "但不同目录是否一定属于不同真实患者尚未完成最终确认，因此目前不能写成"
        "“已经完成患者级独立性验证”。",
    )

    add_heading(doc, "四、在线增强与软件验证", 1)
    add_body(
        doc,
        "首版增强采用保守几何方案：两种模态均以0.5概率水平翻转，以0.5概率在±10°内旋转；"
        "旋转时保留边缘目标，框同步变换。白光另设可选共同增益，荧光共同增益默认关闭。"
        "首版关闭色相、饱和度、随机擦除、随机裁切、噪声、模糊、Mosaic、MixUp和CutMix，"
        "避免小病灶被裁掉后标签仍然保留。",
    )
    for item in [
        "11项合成测试通过，覆盖翻转、边缘旋转、相交框、类别保留、饱和跳过和固定种子重放。",
        "全部1117张图片和1200个框完成加载检查，验证框数量、类别顺序和变换后坐标。",
        "验证集与测试集重复读取结果一致，4485项包内SHA-256校验通过。",
        "尚未完成新版临床图像上的正式训练、增强方案消融或模型性能评估。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "五、模型训练探索", 1)
    add_heading(doc, "5.1 旧版590张数据上的二分类探索", 2)
    add_body(
        doc,
        "为使流程尽早跑通，模型组先在旧版590张荧光数据上完成了二分类训练、验证集阈值选择、"
        "患者级评价和过程可视化。所有结果均为旧版数据、单次患者切分，不能代表新版数据性能。",
    )
    add_table(
        doc,
        ["实验", "患者AUROC", "选定阈值", "患者Balanced Acc"],
        [
            ["ResNet18，全图继承患者标签", "0.639", "0.15", "0.528"],
            ["ResNet18，仅有效框图片", "0.722", "0.39", "0.597"],
            ["ResNet18，冻结主干线性探针", "0.583", "0.26", "0.542"],
            ["颜色/纹理逻辑回归", "0.222", "0.46", "0.222"],
        ],
        widths=[8.2, 3.0, 2.8, 3.8],
    )
    add_callout(
        doc,
        "历史编码说明",
        "旧版探索实验训练时将 n 作为类别1、l 作为类别0。表中 AUROC 不受阳性类名称交换影响，"
        "但选定阈值、敏感度、特异度和部分 Balanced Acc 解释需要按修正后的 n=0、l=1 重新计算。"
        "旧结果只保留为历史参考。",
        "FFF2CC",
    )
    add_body(
        doc,
        "当前最佳患者AUROC约0.72，但测试集只有13位患者、4位阳性，bootstrap 95% CI约为"
        "0.33至1.00；阈值从0.39至0.63变化时，敏感度和特异度也随之明显变化。"
        "结果尚不足以支持论文主结论或临床应用。",
    )
    add_figure(doc, TRAINING_FIGURE, "图1  旧版590张数据上的ResNet18训练过程（探索性结果）")
    add_figure(doc, TEST_FIGURE, "图2  旧版590张数据上的测试概览（探索性结果）")

    add_heading(doc, "5.2 指导书参数与最优模型选择", 2)
    add_body(
        doc,
        "按照项目指导书，已实现 ResNet-50 基线与 EfficientNet-B3 主模型的统一训练入口，"
        "包括 ImageNet 预训练、Adam、三阶段解冻、学习率1e-3/1e-4/1e-5、20/40/40 epoch、"
        "Dropout 0.5、标签平滑0.1、权重衰减1e-4、patience 10、类别加权和余弦退火。"
        "最优检查点按验证集准确率选择，准确率相同时依次比较macro F1和AUROC。",
    )
    add_callout(
        doc,
        "执行状态",
        "训练框架已完成1 epoch/阶段冒烟验证，ResNet-50 与 EfficientNet-B3 均可正常运行并输出"
        "验证指标；正式20/40/40训练尚未完成，因此当前没有可报告的新版指导书对照结果。",
        "FFF2CC",
    )

    add_heading(doc, "六、当前主要问题", 1)
    add_table(
        doc,
        ["问题", "影响", "当前处理"],
        [
            ["数据版本并存", "旧版590与新版1410不能直接合并或比较", "分版本记录，优先统一到新版包"],
            ["患者关联未完成", "来源组划分不等于真实患者级独立", "继续核对登记表、设备字段和病例说明"],
            ["五分类标签不足", "无法直接执行指导书五分类任务", "先做l/n框级二分类，五分类暂缓"],
            ["模型样本量小", "测试集仅13位患者，置信区间很宽", "增加重复分组切分和外部验证"],
            ["图像级标签噪声", "患者标签继承到所有图片可能造成假阴性", "优先使用框级或逐图确认标签"],
            ["隐私与来源管理", "公开仓库历史中仍有身份信息风险", "私密证据分离，清理公开提交历史"],
        ],
        widths=[4.4, 7.0, 5.2],
    )

    add_heading(doc, "七、下一阶段计划", 1)
    add_table(
        doc,
        ["优先级", "任务", "完成标准"],
        [
            ["P0", "冻结新版数据包与类别字典", "图片、XML、TXT、划分、校验和和版本说明一致"],
            ["P0", "完成患者关联核对", "明确同一患者跨目录关系，重新检查跨集风险"],
            ["P0", "固化框级标签规则", "形成统一的数据字典与排除标准"],
            ["P1", "在新版数据上执行指导书训练", "ResNet-50与EfficientNet-B3完成正式训练和验证选择"],
            ["P1", "建立重复分组验证", "多次来源组或患者级切分，报告均值、标准差和95% CI"],
            ["P1", "统一评价体系", "Accuracy、macro F1、AUROC、Kappa、敏感度、特异度与校准"],
            ["P2", "增强消融", "固定预处理、仅几何增强、单项光度增强逐项比较"],
            ["P2", "外部或时间外验证", "至少一个独立数据版本或时间批次"],
        ],
        widths=[2.6, 7.6, 6.4],
    )

    add_heading(doc, "7.1 未来预期：五分类标注到位后的模型迁移", 2)
    add_body(
        doc,
        "如果后续获得经过临床确认的五分类标注，现有模型框架具备较快迁移的基础。"
        "ResNet-50或EfficientNet-B3主干、ImageNet预训练、三阶段解冻流程、数据增强、"
        "训练入口和最优检查点保存机制都可以继续复用，主要改动集中在类别映射、分类头、"
        "类别权重和评价指标。",
    )
    add_table(
        doc,
        ["五分类标注粒度", "可复用部分", "仍需改造", "迁移预期"],
        [
            [
                "整图或患者级五分类",
                "主干网络、迁移学习、训练流程、数据划分",
                "输出层2类改5类、类别权重、macro F1/AUC/混淆矩阵",
                "工程改动较小，标注确认和重新训练是主要耗时",
            ],
            [
                "框级五分类",
                "主干特征、数据增强、基础训练与校验代码",
                "检测/ROI流程、框级类别映射、NMS与FROC评价",
                "需要新增框级分类或五类检测流程，不能只替换最后全连接层",
            ],
            [
                "患者级与框级标签混合",
                "现有数据治理、切分和基础模型",
                "先明确预测单位和金标准，再设计多任务或分层模型",
                "迁移前必须完成任务定义，不能直接开始训练",
            ],
        ],
        widths=[4.0, 5.4, 6.0, 6.0],
    )
    add_body(
        doc,
        "预期结论是“有条件地快速适配”：如果五分类标签以整图或患者为单位、类别定义清楚且患者切分稳定，"
        "模型头从2类替换为5类并重新训练即可形成首版五分类模型；主要时间将花在标签核验和重复验证上。"
        "如果最终目标是框级五分类，现有二分类分类头不能直接解决，需要在检测模型或ROI分类器上继续开发。"
        "因此不能把迁移简单理解为改一个输出层，数据粒度和评价单位仍是决定性因素。",
    )

    add_heading(doc, "八、结论", 1)
    add_body(
        doc,
        "本阶段已经完成从原始数据盘点到可交付二分类候选包的主体工作，并建立了可运行的"
        "模型训练、评价和可视化流程。当前最重要的不是继续堆叠模型结构，而是完成新版数据的"
        "患者关联、患者级独立切分和正式训练。只有在统一数据版本和评价口径后，"
        "ResNet-50与EfficientNet-B3的比较才具有论文意义。",
    )

    add_heading(doc, "资料来源与说明", 1)
    for item in [
        "陈明睿：《口腔荧光图像数据处理阶段工作报告》，2026-10-19。",
        "项目仓库：docs/experiments/OR-B-261019-001.md、OR-B-261019-002.md、OR-B-261019-003.md。",
        "项目仓库：configs/binary_plan.yaml、scripts/train_binary_plan.py。",
        "项目仓库：docs/会议记录.md、docs/论文进度.md。",
    ]:
        add_bullet(doc, item)
    add_body(
        doc,
        "说明：本报告为阶段进展总结，不替代原始实验日志、数据交接单或临床标注确认文件。"
        "报告中的模型指标只对应旧版590张数据，不可与新版1117张候选数据直接比较。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"saved={OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    build_report()

