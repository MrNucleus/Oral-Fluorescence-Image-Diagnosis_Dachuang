"""Generate a training-focused project progress report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from generate_progress_report import (
    DARK,
    GRAY,
    NAVY,
    PROJECT_ROOT,
    add_body,
    add_bullet,
    add_callout,
    add_figure,
    add_heading,
    add_table,
    configure_document,
    set_run_font,
)


OUTPUT = (
    PROJECT_ROOT
    / "docs"
    / "口腔荧光图像AI辅助诊断_模型训练阶段工作报告_2026-10-19_严子宁.docx"
)

TRAINING_FIGURE = (
    PROJECT_ROOT
    / "experiments"
    / "binary_resnet18_annotated_seed43"
    / "training_curves.png"
)
TEST_FIGURE = (
    PROJECT_ROOT
    / "experiments"
    / "binary_resnet18_annotated_seed43"
    / "test_overview.png"
)


def build_report():
    doc = Document()
    configure_document(doc)
    doc.sections[0].header.paragraphs[0].text = ""
    header_run = doc.sections[0].header.paragraphs[0].add_run(
        "口腔荧光图像AI辅助诊断｜模型训练阶段工作报告"
    )
    set_run_font(header_run, "微软雅黑", 8.5, False, GRAY)
    doc.core_properties.title = "口腔荧光图像AI辅助诊断模型训练阶段工作报告"
    doc.core_properties.subject = "模型训练阶段工作总结"
    doc.core_properties.author = "严子宁"

    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("口腔荧光图像AI辅助诊断\n模型训练阶段工作报告")
    set_run_font(run, "微软雅黑", 24, True, NAVY)
    title.paragraph_format.space_after = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("二分类实验、评价流程与指导书训练框架")
    set_run_font(run, "微软雅黑", 13.5, False, GRAY)

    for _ in range(5):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("汇报人：严子宁\n2026年10月19日")
    set_run_font(run, "宋体", 12, False, DARK)
    meta.paragraph_format.line_spacing = 1.8
    doc.add_page_break()

    add_heading(doc, "一、本人职责与阶段结论", 1)
    add_callout(
        doc,
        "当前结论",
        "本人负责模型训练、评估、最优模型选择、实验记录和过程可视化。当前已经完成"
        "二分类训练流程、患者级评价和指导书三阶段训练框架，但尚未得到稳定、可用于"
        "论文主结论的模型。旧版590张数据上的探索性患者AUROC最高约0.72，测试集只有"
        "13位患者，结果不能外推到新版数据或临床场景。",
    )
    add_body(
        doc,
        "本报告只写模型训练相关工作。数据清洗、标注整理、脱敏和来源组划分由数据组完成，"
        "在本报告中只作为训练输入和版本约束出现。",
    )

    add_heading(doc, "二、训练输入与标签口径", 1)
    add_body(
        doc,
        "当前仓库中用于模型训练的主要是旧版590张荧光图像，包含510份XML标注，按患者"
        "整理为83个患者标识。仅使用有有效框的图片时，可用样本为509张，训练、验证、"
        "测试分别为339、88、82张，对应57、13、13位患者。",
    )
    add_callout(
        doc,
        "标签定义",
        "l = 阳性病例，编码1；n = 阴性病例，编码0。新生成的binary_v2和"
        "binary_annotated_v2清单已经按这一编码固定。",
    )
    add_body(
        doc,
        "早期模型实验使用了相反的编码，即n=1、l=0。因此旧实验中的敏感度、精确率、"
        "召回率和验证集选定阈值均不能直接作为l阳性指标复用；AUROC在二分类中不受阳性类"
        "交换影响，但阈值相关结果必须重新计算。",
    )
    add_table(
        doc,
        ["输入项目", "旧版探索数据", "状态"],
        [
            ["图像总数", "590", "已用于探索"],
            ["XML数", "510", "已配对"],
            ["患者标识", "83", "已有患者级切分"],
            ["有有效框图片", "509", "当前分类实验使用"],
            ["训练/验证/测试图片", "339/88/82", "固定切分"],
            ["训练/验证/测试患者", "57/13/13", "无患者重叠"],
            ["新版v5候选包", "1117张、1200框", "尚未接入当前训练"],
        ],
        widths=[5.6, 5.4, 5.6],
    )

    add_heading(doc, "三、已完成的训练流程", 1)
    add_heading(doc, "3.1 数据与评价流程", 2)
    for item in [
        "建立匿名患者清单，使用患者级切分，避免同一患者跨训练、验证和测试集。",
        "实现全图缩放、水平翻转、轻度旋转、亮度和对比度调整，不进行随机裁切，避免直接裁掉病灶。",
        "实现类别加权交叉熵、AMP混合精度、AdamW优化器和余弦退火。",
        "记录训练损失、训练准确率、验证图像级指标、验证患者级指标和每轮耗时。",
        "实现验证集阈值选择、患者级均值聚合和患者cluster bootstrap置信区间。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.2 ResNet18二分类实验", 2)
    add_body(
        doc,
        "先使用普通ResNet18 ImageNet预训练模型建立基线。第一阶段冻结主干，只训练分类头；"
        "第二阶段解冻layer4进行微调。模型选择使用验证集患者级AUROC，测试集只在模型"
        "选择完成后评估。",
    )
    add_table(
        doc,
        ["实验", "测试患者AUROC", "选定阈值", "患者Balanced Acc", "主要判断"],
        [
            ["全图继承患者标签", "0.639", "0.15", "0.528", "旧编码，结果不稳定"],
            ["仅有效框图片", "0.722", "0.39", "0.597", "当前最高，但仍不可用"],
            ["冻结主干线性探针", "0.583", "0.26", "0.542", "降低自由度没有解决问题"],
            ["颜色/纹理逻辑回归", "0.222", "0.46", "0.222", "简单颜色信号不稳定"],
        ],
        widths=[5.0, 3.0, 2.6, 3.4, 5.4],
    )
    add_body(
        doc,
        "微调阶段训练准确率快速升高，但验证集患者AUROC没有同步改善，出现明显过拟合。"
        "冻结主干后测试AUROC反而下降，说明问题不只是模型容量，可能还包括图像级标签噪声、"
        "患者间采集差异和有效样本量不足。",
    )

    add_heading(doc, "3.3 过程可视化", 2)
    add_body(
        doc,
        "为每次实验生成了训练曲线、训练过程GIF、测试概览和HTML看板。可视化包括训练与验证"
        "指标、患者概率分布、图像级ROC、图像级和患者级混淆矩阵，便于检查模型是在学习还是"
        "只是在拟合训练集。",
    )
    add_figure(doc, TRAINING_FIGURE, "图1  旧版509张数据上的ResNet18训练过程（历史编码）")
    add_figure(doc, TEST_FIGURE, "图2  旧版509张数据上的测试概览（历史编码）")

    add_heading(doc, "四、按照指导书实现的模型参数与最优模型选择", 1)
    add_body(
        doc,
        "已实现ResNet-50基线和EfficientNet-B3主模型的统一训练入口。参数按项目指导书设置："
        "ImageNet预训练、Adam、三阶段解冻、学习率1e-3/1e-4/1e-5、20/40/40个epoch、"
        "Dropout 0.5、标签平滑0.1、权重衰减1e-4和早停patience 10。",
    )
    add_table(
        doc,
        ["项目", "实现内容"],
        [
            ["基线模型", "ResNet-50"],
            ["主要优化模型", "EfficientNet-B3"],
            ["优化器", "Adam"],
            ["阶段1", "冻结主干，只训练分类头，lr=1e-3，20 epoch"],
            ["阶段2", "解冻最后两个stage，lr=1e-4，40 epoch"],
            ["阶段3", "解冻全部参数，lr=1e-5，40 epoch"],
            ["分类头", "Dropout 0.5 + Linear + ReLU + Dropout 0.25 + Linear"],
            ["损失函数", "类别加权交叉熵 + label smoothing 0.1"],
            ["最优模型选择", "验证集准确率；同分依次比较macro F1和AUROC"],
        ],
        widths=[5.0, 11.0],
    )
    add_callout(
        doc,
        "执行状态",
        "训练框架已经完成1 epoch/阶段的冒烟验证，ResNet-50和EfficientNet-B3均能正常"
        "运行并保存验证结果；正式20/40/40训练尚未完成。当前没有可报告的新版指导书对照结果。",
        "FFF2CC",
    )
    add_body(
        doc,
        "由于当前只有二分类标签，本轮严格照搬了指导书的模型结构、优化器、阶段学习率、正则化"
        "和最优模型选择规则，但没有执行五分类类别数和7:2:1划分。现有患者级70/15/15切分"
        "继续保留，以保证旧实验可比。",
    )

    add_heading(doc, "五、当前主要问题", 1)
    add_table(
        doc,
        ["问题", "对训练的影响", "下一步"],
        [
            ["有效样本量小", "测试集仅13位患者，指标波动大", "重复患者级切分或分组交叉验证"],
            ["图像级标签可能含噪", "患者标签继承给所有图片，可能把无病灶图当作阳性", "改为逐图标签或框级ROI"],
            ["旧版与新版数据不一致", "当前指标不能代表v5数据", "将训练切换至新版1117张候选包"],
            ["旧实验标签编码相反", "旧敏感度和阈值不能直接复用", "使用l=1、n=0的v2清单重新训练"],
            ["微调易过拟合", "训练准确率接近饱和，验证AUROC下降", "优先线性探针、强正则和冻结策略"],
            ["正式指导书训练未跑完", "尚无ResNet-50与EfficientNet-B3正式比较", "完成20/40/40训练和验证选择"],
            ["缺少外部验证", "无法判断临床泛化", "增加时间外或第二中心验证"],
        ],
        widths=[5.0, 7.8, 6.2],
    )

    add_heading(doc, "六、未来预期：五分类标注到位后的适配", 1)
    add_callout(
        doc,
        "结论",
        "可以在模型结构上较快适配，但能否称为“快速完成”取决于五分类标签粒度、标签质量和"
        "患者级切分，而不是只取决于网络代码。",
    )
    add_body(
        doc,
        "如果五分类标签是整图或患者级，现有ResNet-50/EfficientNet-B3主干、三阶段训练、"
        "数据增强、检查点保存和验证选择流程都可以复用。主要改造是输出层从2类改为5类、"
        "重新计算类别权重、增加macro F1/混淆矩阵/Kappa等指标，然后重新训练。",
    )
    add_body(
        doc,
        "如果五分类标签是框级，则不能只替换最后分类层。需要在检测模型或ROI分类器中增加"
        "五类病种输出，并调整框级类别映射、NMS和FROC评价。此时主要工作量在任务重构和"
        "标签核验，而不是分类头本身。",
    )
    add_table(
        doc,
        ["标注粒度", "迁移预期"],
        [
            ["整图或患者级五分类", "较快：主干和训练流程复用，替换2类头为5类头"],
            ["框级五分类", "中等：需要检测/ROI流程和新的框级评价"],
            ["混合粒度", "先统一标签和预测单位，再决定多任务或分层模型"],
        ],
        widths=[6.0, 10.0],
    )

    add_heading(doc, "七、下一阶段训练计划", 1)
    add_table(
        doc,
        ["优先级", "任务", "验收标准"],
        [
            ["P0", "接入新版v5候选包", "固定图片、XML、TXT和来源组划分"],
            ["P0", "切换到v2标签编码", "全部训练统一使用n=0、l=1"],
            ["P0", "完成ResNet-50正式训练", "按20/40/40跑完并保存最佳检查点"],
            ["P0", "完成EfficientNet-B3正式训练", "按同一规则完成并输出验证选择结果"],
            ["P1", "比较最优模型", "按验证准确率选择，记录F1、AUROC和Kappa"],
            ["P1", "重复患者级验证", "多个种子或分组交叉验证，报告均值和95% CI"],
            ["P1", "增强与预处理消融", "固定数据版本，分别比较原始、几何增强和光度增强"],
            ["P2", "五分类适配", "标签确认后替换输出头并重新训练"],
        ],
        widths=[2.5, 7.3, 7.2],
    )

    add_heading(doc, "八、结论", 1)
    add_body(
        doc,
        "本阶段已经完成从二分类基线到指导书三阶段训练框架的搭建，训练、验证、阈值选择、"
        "患者级评价和可视化流程均可运行。当前主要障碍不是代码能否训练，而是样本量、标签质量、"
        "数据版本和评价口径。下一步应先将训练切换到新版v5数据和v2标签编码，再完成"
        "ResNet-50/EfficientNet-B3正式训练，之后才适合讨论五分类迁移和论文主结果。",
    )

    add_heading(doc, "资料来源", 1)
    add_bullet(doc, "docs/experiments/OR-B-261019-001.md、OR-B-261019-002.md、OR-B-261019-003.md。")
    add_bullet(doc, "configs/binary_plan.yaml、scripts/train_binary_plan.py。")
    add_bullet(doc, "docs/项目全过程日志.md、docs/技术方案.md。")
    add_bullet(doc, "新版数据包的交接信息仅作为训练输入背景引用，数据处理工作不作为本人阶段成果。")

    doc.save(OUTPUT)
    print(f"saved={OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    build_report()

